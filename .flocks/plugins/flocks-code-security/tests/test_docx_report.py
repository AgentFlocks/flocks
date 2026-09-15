from io import BytesIO
from types import SimpleNamespace
import hashlib

import pytest
from docx import Document

from flocks_code_security.docx_report import render_docx
from flocks_code_security.service import AuditService, AuditServiceError


def test_docx_preserves_report_structure_and_code():
    data = render_docx(
        "# 安全审计\n\n- Result: **confirmed**\n\n| File | Risk |\n| --- | --- |\n| api.py | high |\n\n## Evidence\n```python\nif a < b:\n    print('证据 & text')\n```\n"
    )
    document = Document(BytesIO(data))
    assert document.paragraphs[0].text == "安全审计"
    assert document.paragraphs[0].style.name == "Title"
    assert document.tables[0].cell(1, 0).text == "api.py"
    assert "    print('证据 & text')" in [p.text for p in document.paragraphs]
    assert any(run.bold for p in document.paragraphs for run in p.runs if run.text == "confirmed")


@pytest.mark.asyncio
async def test_docx_download_uses_the_existing_access_and_integrity_checks(tmp_path, monkeypatch):
    path = tmp_path / "report.md"
    path.write_text("# Trusted report", encoding="utf-8")
    service = object.__new__(AuditService)
    calls = []
    scan = {"scan_id": "scan", "status": "completed"}

    def visible(scan_id, caller):
        calls.append((scan_id, caller))
        return scan

    service._require_visible_scan = visible
    service._artifact_file = lambda scan_id, kind: path if kind == "report_markdown" else None
    service.store = SimpleNamespace(
        scan_status=lambda _: {
            "integrity_status": "valid",
            "integrity_artifacts": {"report.md": hashlib.sha256(b"# Trusted report").hexdigest()},
        }
    )
    filename, data = await service.download_artifact("scan", "report.docx", "owner")
    assert calls == [("scan", "owner")]
    assert filename == "report.docx"
    assert Document(BytesIO(data)).paragraphs[0].text == "Trusted report"
    import builtins

    original_import = builtins.__import__

    def unavailable(name, *args, **kwargs):
        if name == "flocks_code_security.docx_report":
            raise ModuleNotFoundError("No module named 'docx'", name="docx")
        return original_import(name, *args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(builtins, "__import__", unavailable)
        with pytest.raises(AuditServiceError) as missing:
            await service.download_artifact("scan", "report.docx", "owner")
        assert missing.value.code == "docx_export_unavailable"
        assert missing.value.status_code == 503
    path.write_text("tampered", encoding="utf-8")
    with pytest.raises(AuditServiceError) as exc:
        await service.download_artifact("scan", "report.docx", "owner")
    assert exc.value.code == "artifact_integrity_invalid"


def test_markdown_escapes_links_and_nested_formatting():
    report = r"""# Report

Text with file\_name, \(argument\), **bold and *italic*** and [link](https://example.com).

- Parent
  - Child

1. First
2. Second

| Path | Value |
| --- | --- |
| `a\|b` | **yes** |

````python
value = "```"
path = r"C:\temp\_name"
````
"""
    document = Document(BytesIO(render_docx(report)))
    texts = [p.text for p in document.paragraphs]
    assert any("file_name, (argument)" in text for text in texts)
    assert any(run.bold and run.italic for p in document.paragraphs for run in p.runs)
    assert 'value = "```"' in texts
    assert 'path = r"C:\\temp\\_name"' in texts
    assert document.tables[0].cell(1, 0).text == "a|b"
    assert document.part.rels.values()
    assert any(rel.target_ref == "https://example.com" for rel in document.part.rels.values() if rel.is_external)
    assert any(p.style.name == "List Number" for p in document.paragraphs)


def test_report_inline_escapes_are_removed_without_altering_source_code():
    document = Document(BytesIO(render_docx("`scan\\_123`\n\n```\nvalue = r'file\\_name'\n```", report_escapes=True)))
    assert document.paragraphs[0].text == "scan_123"
    assert "value = r'file\\_name'" in [p.text for p in document.paragraphs]
