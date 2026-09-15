"""Render the verified Markdown report as native Word paragraphs and tables."""

from __future__ import annotations

from io import BytesIO
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt


def _inline(paragraph, tokens, *, report_escapes: bool) -> None:
    bold = italic = strike = False
    link = None
    for token in tokens or []:
        if token.type in {"strong_open", "strong_close"}:
            bold = token.type == "strong_open"
        elif token.type in {"em_open", "em_close"}:
            italic = token.type == "em_open"
        elif token.type in {"s_open", "s_close"}:
            strike = token.type == "s_open"
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type in {"text", "code_inline", "image"}:
            text = token.content
            # The report writer escapes paths even inside inline code. Undo only
            # its known escapes; fenced source code is never passed through here.
            if token.type == "code_inline" and report_escapes:
                text = re.sub(r"\\([\\`*_{}\[\]()<>#!|])", r"\1", text)
            run = paragraph.add_run(text)
            run.bold = True if bold else None
            run.italic = True if italic else None
            run.font.strike = True if strike else None
            if token.type == "code_inline":
                run.font.name = "Courier New"
                run.font.size = Pt(9)
        elif token.type == "link_open":
            # Word hyperlink relationships are deliberately restricted to URLs.
            href = token.attrGet("href") or ""
            if href.startswith(("https://", "http://", "mailto:")):
                from docx.opc.constants import RELATIONSHIP_TYPE

                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                paragraph._p.append(link)
                # Subsequent runs are moved to this link on link_close.
        elif token.type == "link_close":
            if link is not None:
                while link.getnext() is not None:
                    link.append(link.getnext())
                link = None


def render_docx(markdown: str, *, report_escapes: bool = False) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    for name in ["Normal", "Title", *[f"Heading {i}" for i in range(1, 7)]]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    normal = document.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    # XML cannot represent these controls. Do not strip tabs or newlines.
    markdown = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", markdown)
    tokens = MarkdownIt("commonmark", {"html": False}).enable(["table", "strikethrough"]).parse(markdown)
    paragraph = None
    table = row = None
    lists = []
    for token in tokens:
        kind = token.type
        if kind in {"bullet_list_open", "ordered_list_open"}:
            number_id = None
            if kind == "ordered_list_open":
                numbering = document.part.numbering_part.element
                abstract = numbering.xpath("./w:abstractNum[w:lvl/w:pStyle[@w:val='ListNumber']]")[0]
                number = numbering.add_num(int(abstract.get(qn("w:abstractNumId"))))
                number.add_lvlOverride(0).add_startOverride(int(token.attrGet("start") or 1))
                number_id = number.numId
            lists.append({"number_id": number_id, "first": False})
        elif kind == "list_item_open":
            lists[-1]["first"] = True
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif kind == "heading_open":
            paragraph = document.add_paragraph(
                style="Title" if token.tag == "h1" else f"Heading {int(token.tag[1]) - 1}"
            )
        elif kind == "paragraph_open":
            paragraph = document.add_paragraph()
            if lists:
                level = lists[-1]
                paragraph.paragraph_format.left_indent = Inches(0.18 * len(lists))
                if level["first"]:
                    paragraph.style = "List Number" if level["number_id"] is not None else "List Bullet"
                    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
                    if level["number_id"] is not None:
                        properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                        properties.get_or_add_numId().val = level["number_id"]
                        properties.get_or_add_ilvl().val = 0
                    level["first"] = False
        elif kind == "table_open":
            table = document.add_table(rows=0, cols=0)
            table.style = "Light Shading Accent 1"
        elif kind == "tr_open":
            row = table.add_row()
            column = 0
        elif kind in {"th_open", "td_open"}:
            if column >= len(table.columns):
                table.add_column(Inches(1))
            cell = row.cells[column]
            column += 1
            paragraph = cell.paragraphs[0]
            if kind == "th_open":
                paragraph.style = "Normal"
                props = row._tr.get_or_add_trPr()
                if props.find(qn("w:tblHeader")) is None:
                    props.append(OxmlElement("w:tblHeader"))
        elif kind == "inline" and paragraph is not None:
            _inline(paragraph, token.children, report_escapes=report_escapes)
        elif kind == "table_close":
            width = (section.page_width - section.left_margin - section.right_margin) // len(table.columns)
            for column in table.columns:
                column.width = width
            for table_row in table.rows:
                for table_cell in table_row.cells:
                    table_cell.width = width
            table = row = None
        elif kind in {"fence", "code_block"}:
            # Keep each source line independent so long excerpts can paginate.
            for line in token.content.rstrip("\n").split("\n"):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                paragraph.paragraph_format.widow_control = False
                run = paragraph.add_run(line)
                run.font.name, run.font.size = "Courier New", Pt(8)
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F3F4F6")
                paragraph._p.get_or_add_pPr().append(shade)
            document.add_paragraph().paragraph_format.space_after = Pt(0)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
